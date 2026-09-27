from datetime import datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from app.database.models import Order, Stock, User
from app.database.repositories.cart import CartRepository
from app.database.repositories.orders import OrderRepository
from app.database.repositories.stats import StatsRepository, period_start
from app.services.stats import orders_excel, stats_text
from tests.test_offers_cart import db  # noqa: F401

TZ = ZoneInfo("Asia/Tashkent")


def test_period_start_uses_local_midnight():
    now = datetime(2026, 9, 27, 18, 30, tzinfo=timezone.utc)  # 23:30 в Ташкенте
    assert period_start("today", TZ, now) == datetime(2026, 9, 26, 19, 0, tzinfo=timezone.utc)  # 00:00 27.09 по Ташкенту
    assert period_start("7", TZ, now) == datetime(2026, 9, 20, 19, 0, tzinfo=timezone.utc)
    assert period_start("all", TZ, now) is None
    late = datetime(2026, 9, 27, 19, 30, tzinfo=timezone.utc)  # 00:30 28.09 — уже новый день
    assert period_start("today", TZ, late) == datetime(2026, 9, 27, 19, 0, tzinfo=timezone.utc)


async def make_order(db, user, stock_id, qty=1, days_ago=0, status="NEW"):
    cart = CartRepository(db)
    for _ in range(qty):
        await cart.add(user, await db.get(Stock, stock_id))
    order = (await OrderRepository(db).create_from_cart(user))[0]
    order.created_at = datetime.now(timezone.utc) - timedelta(days=days_ago)
    order.status = status
    await db.commit()
    return order


async def test_stats_by_period(db):
    user = await db.get(User, 1)
    user.phone = "+998901234567"
    await make_order(db, user, 1, qty=2)  # сегодня, 170 000, дилер №1
    await make_order(db, user, 2, qty=1, status="COMPLETED")  # сегодня, 82 000
    await make_order(db, user, 1, qty=1, status="CANCELLED")  # сегодня, отменён — не в сумме
    await make_order(db, user, 2, qty=1, days_ago=10)  # 10 дней назад

    repo = StatsRepository(db)
    today = await repo.collect(period_start("today", TZ))
    assert today.orders == 3 and today.revenue == Decimal("252000") and today.completed_sum == Decimal("82000")
    assert today.avg_check == Decimal("126000") and today.by_status == {"NEW": 1, "COMPLETED": 1, "CANCELLED": 1}
    assert [(d, n) for d, n, _ in today.by_dealer] == [("JAC Ташкент №1", 1), ("JAC Ташкент №2", 1)]
    assert today.top_parts == [("Масляный фильтр", "101001", 3)]
    assert today.by_region[0][0].name_ru == "Ташкент" and today.active_clients == 1

    month = await repo.collect(period_start("30", TZ))
    assert month.orders == 4 and month.revenue == Decimal("334000")

    text = stats_text(today, "today", "ru")
    assert "📊 СТАТИСТИКА · Сегодня" in text and "🧾 Заказов: 3" in text and "252 000 сум" in text
    assert "1. Масляный фильтр (101001) — 3 шт." in text


async def test_excel_export(db):
    user = await db.get(User, 1)
    user.phone = "+998901234567"
    await make_order(db, user, 1, qty=2)
    orders = await StatsRepository(db).orders_for_export(None)
    for o in orders:
        await db.refresh(o)
    wb = load_workbook(BytesIO(orders_excel(orders, "ru")))
    assert wb.sheetnames == ["Заказы", "Позиции"]
    head, row = [c.value for c in wb["Заказы"][1]], [c.value for c in wb["Заказы"][2]]
    assert head[0] == "№ заказа" and row[0] == orders[0].id and row[4] == "+998901234567" and row[9] == 170000
    item = [c.value for c in wb["Позиции"][2]]
    assert item[4] == "Масляный фильтр" and item[5] == "101001" and item[8] == 2


async def test_empty_stats(db):
    s = await StatsRepository(db).collect(None)
    assert s.orders == 0 and "За этот период заказов нет." in stats_text(s, "all", "ru")
    assert isinstance(s, type(s)) and Order  # импорт используется
