"""Файлы в чате: что как показывать прямо в переписке и как отдавать их браузеру.

Telegram хранит файл у себя, в базе — только file_id. При первом просмотре панель скачивает файл через бота
(до 20 МБ — ограничение Telegram для ботов) и кладёт во временную папку: повторный просмотр и перемотка видео
не скачивают его заново. Файлы старше недели удаляются сами.

Показываем в самом чате: фото и картинки, видео, голосовые и аудио, PDF; таблицы Excel/CSV, документы Word и
текстовые файлы — предпросмотром. Опасные для браузера форматы (HTML, SVG, скрипты…) открыть нельзя — только скачать.
"""
from __future__ import annotations

import csv
import hashlib
import io
import os
import time
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree

from aiogram import Bot

from app.database.models import SupportMessage

CACHE_DIR = Path(os.getenv("CHAT_CACHE_DIR", "/tmp/jac_chat_files"))
KEEP_SECONDS = 7 * 24 * 3600
MAX_ROWS, MAX_COLS, MAX_TEXT = 200, 30, 200_000

EXT = {
    "image": ({"jpg", "jpeg", "png", "webp", "gif"}, None),
    "video": ({"mp4", "m4v", "mov", "webm"}, None),
    "audio": ({"mp3", "ogg", "oga", "opus", "m4a", "wav"}, None),
    "pdf": ({"pdf"}, "application/pdf"),
    "table": ({"xlsx", "xlsm", "csv"}, None),
    "word": ({"docx"}, None),
    "text": ({"txt", "log", "json", "xml", "md"}, "text/plain; charset=utf-8"),
}
MIME = {"jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png", "webp": "image/webp", "gif": "image/gif",
        "mp4": "video/mp4", "m4v": "video/mp4", "mov": "video/quicktime", "webm": "video/webm",
        "mp3": "audio/mpeg", "ogg": "audio/ogg", "oga": "audio/ogg", "opus": "audio/ogg", "m4a": "audio/mp4",
        "wav": "audio/wav"}
ICONS = {"pdf": "📕", "table": "📊", "word": "📘", "text": "📄", "video": "🎬", "audio": "🎵", "image": "🖼",
         "other": "📎"}


@dataclass
class View:
    kind: str  # image | video | audio | pdf | table | word | text | other
    mime: str
    inline: bool  # браузер может безопасно показать сам

    @property
    def icon(self) -> str:
        return ICONS.get(self.kind, "📎")

    @property
    def previewable(self) -> bool:
        return self.kind in ("pdf", "table", "word", "text", "image")


def _ext(name: str | None) -> str:
    return (name or "").rsplit(".", 1)[-1].lower() if "." in (name or "") else ""


def view_of(msg: SupportMessage) -> View:
    """Как показывать файл сообщения."""
    kind = msg.file_kind
    if kind == "photo":
        return View("image", "image/jpeg", True)
    if kind == "sticker":
        return View("image", "image/webp", True)
    if kind == "voice":
        return View("audio", "audio/ogg", True)
    if kind in ("video", "video_note"):
        return View("video", MIME.get(_ext(msg.file_name), "video/mp4"), True)
    ext = _ext(msg.file_name)
    for view_kind, (exts, mime) in EXT.items():
        if ext in exts:
            mime = mime or MIME.get(ext, "application/octet-stream")
            inline = view_kind in ("image", "video", "audio", "pdf", "text")
            return View(view_kind, mime, inline)
    if kind == "audio":
        return View("audio", "audio/mpeg", True)
    return View("other", "application/octet-stream", False)


def _cleanup() -> None:
    now = time.time()
    for path in CACHE_DIR.glob("*"):
        try:
            if now - path.stat().st_mtime > KEEP_SECONDS:
                path.unlink()
        except OSError:
            pass


async def local_copy(bot: Bot, msg: SupportMessage) -> Path:
    """Файл сообщения на диске (скачивается у Telegram один раз). Ошибки Telegram пробрасываются."""
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    path = CACHE_DIR / f"msg_{msg.id}_{hashlib.sha1(msg.file_id.encode()).hexdigest()[:10]}"
    if path.exists() and path.stat().st_size > 0:
        return path
    tg_file = await bot.get_file(msg.file_id)
    data = await bot.download_file(tg_file.file_path)
    content = data.read() if hasattr(data, "read") else data
    tmp = path.with_name(f"{path.name}.{os.getpid()}.part")
    tmp.write_bytes(content)
    tmp.replace(path)  # атомарно: два одновременных запроса не увидят половину файла
    _cleanup()
    return path


# ---------- предпросмотр документов (таблица / Word / текст) ----------

@dataclass
class Preview:
    kind: str  # table | text | error
    sheets: list[tuple[str, list[list[str]], bool]] | None = None  # (лист, строки, обрезано)
    text: str = ""
    error: str = ""


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value)


def _table(path: Path, ext: str) -> Preview:
    if ext == "csv":
        raw = path.read_bytes()
        text = next((raw.decode(enc) for enc in ("utf-8-sig", "cp1251") if _decodes(raw, enc)), raw.decode("latin-1"))
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t") if text.strip() else csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
        return Preview("table", sheets=[("CSV", [r[:MAX_COLS] for r in rows[:MAX_ROWS]], len(rows) > MAX_ROWS)])
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=True)  # в кэше файл без .xlsx
    sheets = []
    for ws in book.worksheets[:5]:
        rows, cut = [], False
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_ROWS:
                cut = True
                break
            rows.append([_cell(v) for v in row[:MAX_COLS]])
        while rows and not any(rows[-1]):  # пустой «хвост» листа
            rows.pop()
        sheets.append((ws.title, rows, cut))
    book.close()
    return Preview("table", sheets=sheets)


def _decodes(raw: bytes, enc: str) -> bool:
    try:
        raw.decode(enc)
        return True
    except UnicodeDecodeError:
        return False


def _word(path: Path) -> Preview:
    """Текст абзацев .docx (это zip с word/document.xml) — без дополнительных библиотек."""
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(path) as z:
        root = ElementTree.fromstring(z.read("word/document.xml"))
    paragraphs = ["".join(t.text or "" for t in p.iter(f"{ns}t")) for p in root.iter(f"{ns}p")]
    return Preview("text", text="\n".join(paragraphs).strip()[:MAX_TEXT])


def preview(path: Path, msg: SupportMessage) -> Preview:
    view, ext = view_of(msg), _ext(msg.file_name)
    try:
        if view.kind == "table":
            return _table(path, ext)
        if view.kind == "word":
            return _word(path)
        if view.kind == "text":
            raw = path.read_bytes()[:MAX_TEXT]
            return Preview("text", text=raw.decode("utf-8", errors="replace"))
    except Exception as e:  # повреждённый или необычный файл — предлагаем скачать
        return Preview("error", error=f"Не удалось показать файл ({type(e).__name__}). Скачайте его.")
    return Preview("error", error="Для этого типа файла предпросмотра нет — скачайте его.")
