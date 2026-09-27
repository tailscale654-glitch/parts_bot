"""Шаблон Excel для загрузки каталога (бот отправляет его по кнопке «📄 Шаблон Excel»)."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.excel_import import OPTIONAL, REQUIRED

COLUMNS = [
    # колонка, пример, описание
    ("region", "Ташкент", "Регион дилера — как в боте (на русском, английском или узбекском)"),
    ("dealer", "JAC Ташкент №1", "Название дилера"),
    ("model_ru", "JAC JS4", "Модель (русский)"),
    ("model_en", "JAC JS4", "Модель (английский) — можно пусто"),
    ("model_uz", "JAC JS4", "Модель (узбекский) — можно пусто"),
    ("node_ru", "Двигатель", "Категория (русский)"),
    ("node_en", "Engine", "Категория (английский) — можно пусто"),
    ("node_uz", "Dvigatel", "Категория (узбекский) — можно пусто"),
    ("part_name_ru", "Масляный фильтр", "Название детали (русский)"),
    ("part_name_en", "Oil filter", "Название детали (английский) — можно пусто"),
    ("part_name_uz", "Moy filtri", "Название детали (узбекский) — можно пусто"),
    ("part_number", "101001", "Артикул. Одинаковый артикул у разных дилеров = одна деталь"),
    ("price", "85000", "Цена в сумах, число больше 0"),
    ("stock", "12", "Остаток, целое число (0 = нет в наличии)"),
    ("delivery_days", "", "Срок доставки в днях, если нет в наличии — можно пусто"),
    ("description_ru", "", "Описание (русский) — можно пусто"),
    ("description_en", "", "Описание (английский) — можно пусто"),
    ("description_uz", "", "Описание (узбекский) — можно пусто"),
    ("photo", "", "Ссылка на фото https://... — можно пусто"),
]

EXAMPLE_ROWS = [
    ["Ташкент", "JAC Ташкент №1", "JAC JS4", "JAC JS4", "JAC JS4", "Двигатель", "Engine", "Dvigatel",
     "Масляный фильтр", "Oil filter", "Moy filtri", "101001", 85000, 12, None],
    ["Ташкент", "JAC Ташкент №2", "JAC JS4", "JAC JS4", "JAC JS4", "Двигатель", "Engine", "Dvigatel",
     "Масляный фильтр", "Oil filter", "Moy filtri", "101001", 82000, 4, None],
    ["Самаркандская область", "JAC Самарканд", "JAC JS4", "JAC JS4", "JAC JS4", "Тормозная система",
     "Brake system", "Tormoz tizimi", "Тормозные колодки передние", "Front brake pads",
     "Old tormoz kolodkalari", "102001", 240000, 0, 3],
]

# Защита от опечаток: шаблон должен совпадать с колонками, которые ждёт импорт
assert set(REQUIRED + OPTIONAL) == {c for c, _, _ in COLUMNS}


def build_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "parts"
    ws.append([c for c, _, _ in COLUMNS])
    for row in EXAMPLE_ROWS:
        ws.append(row)
    required_fill = PatternFill("solid", fgColor="FFE699")
    optional_fill = PatternFill("solid", fgColor="DDEBF7")
    for i, (col, _, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True)
        cell.fill = required_fill if col in REQUIRED else optional_fill
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(col) + 4)
    for row in ws.iter_rows(min_row=2, min_col=12, max_col=12):  # артикул — текст, чтобы не терялись нули
        for cell in row:
            cell.number_format = "@"
    ws.freeze_panes = "A2"

    help_ws = wb.create_sheet("инструкция")
    help_ws.append(["Колонка", "Обязательная", "Что писать"])
    for col, _, desc in COLUMNS:
        help_ws.append([col, "да" if col in REQUIRED else "нет", desc])
    help_ws.append([])
    help_ws.append(["Важно", "", "Одна строка = одна деталь у одного дилера. Лист должен называться parts."])
    help_ws.append(["", "", "Файл — это ПОЛНЫЙ прайс: детали, которых нет в файле, будут скрыты из бота."])
    help_ws.append(["", "", "Жёлтые колонки обязательные, голубые — можно оставить пустыми."])
    for cell in help_ws[1]:
        cell.font = Font(bold=True)
    help_ws.column_dimensions["A"].width = 18
    help_ws.column_dimensions["B"].width = 14
    help_ws.column_dimensions["C"].width = 90
    for row in help_ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
