"""Текст статистики и выгрузка заказов в Excel."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.database.models import Order
from app.database.repositories.stats import Stats
from app.services.catalog import money
from app.services.localization import i18n, localized_name
from app.services.notify import person_name
from app.services import orders as orders_service
from app.services.orders import fmt_date, status_text

STATUS_KEYS = ("NEW", "CONFIRMED", "READY", "COMPLETED", "CANCELLED")


def stats_text(s: Stats, period: str, lang: str | None, recent: list[Order] | None = None) -> str:
    lines = [
        i18n.t(lang, "stats_title", period=i18n.t(lang, f"period_{period}")), "",
        i18n.t(lang, "stats_orders", n=s.orders),
        i18n.t(lang, "stats_revenue", total=money(s.revenue, lang)),
        i18n.t(lang, "stats_completed", total=money(s.completed_sum, lang)),
        i18n.t(lang, "stats_avg", total=money(s.avg_check, lang)),
    ]
    if s.orders:
        lines += [""] + [f"{status_text(k, lang).split(' —')[0]}: {s.by_status.get(k, 0)}" for k in STATUS_KEYS]
    lines += ["", i18n.t(lang, "stats_clients", new=s.new_clients, active=s.active_clients)]
    if s.by_region:
        lines += ["", i18n.t(lang, "stats_by_region")]
        lines += [f"• {localized_name(r, lang)}: {n} · {money(t, lang)}" for r, n, t in s.by_region]
    if s.by_dealer:
        lines += ["", i18n.t(lang, "stats_by_dealer")]
        lines += [f"{i}. {name}: {n} · {money(t, lang)}" for i, (name, n, t) in enumerate(s.by_dealer, 1)]
    if s.top_parts:
        lines += ["", i18n.t(lang, "stats_top_parts")]
        lines += [f"{i}. {name} ({number}) — {qty} {i18n.t(lang, 'pcs')}" for i, (name, number, qty) in enumerate(s.top_parts, 1)]
    if recent:
        lines += ["", i18n.t(lang, "stats_recent")]
        for o in recent:
            lines += [""] + order_with_items(o, lang)
    if not s.orders:
        lines += ["", i18n.t(lang, "stats_no_orders")]
    text = "\n".join(lines)
    if len(text) > 4000:  # лимит Telegram — обрезаем по границе строки
        text = text[:3990].rsplit("\n", 1)[0] + "\n…"
    return text


def order_with_items(o: Order, lang: str | None) -> list[str]:
    """«№10004 · 🟠 Готов к выдаче · 27.09 23:28» + клиент, дилер и каждая позиция."""
    status = status_text(o.status, lang).split(" —")[0]
    lines = [
        f"№{o.id} · {status} · {fmt_date(o.created_at)}",
        f"👤 {person_name(o.user)} · {o.user.phone or '—'}",
        f"🏢 {o.dealer.name}",
    ]
    for item in o.items:
        lines.append(f"   • {localized_name(item, lang)} ({item.part_number}) × {item.quantity} = {money(item.total, lang)}")
    lines.append(f"   💰 {money(o.total_amount, lang)}")
    return lines


def orders_with_items_text(orders: list[Order], lang: str | None, title: str) -> str:
    lines = [title]
    for o in orders:
        lines += [""] + order_with_items(o, lang)
    return "\n".join(lines)


def orders_excel(orders: list[Order], lang: str | None) -> bytes:
    """Два листа: «Заказы» (по строке на заказ) и «Позиции» (по строке на деталь)."""
    wb = Workbook()
    ws = wb.active
    ws.title = i18n.t(lang, "xl_orders_sheet")[:31]
    ws.append([i18n.t(lang, f"xl_{c}") for c in (
        "number", "date", "status", "client", "phone", "client_region", "dealer", "dealer_region", "items", "total")])
    items_ws = wb.create_sheet(i18n.t(lang, "xl_items_sheet")[:31])
    items_ws.append([i18n.t(lang, f"xl_{c}") for c in (
        "number", "date", "status", "client", "phone", "dealer", "part", "part_number", "model", "category", "qty", "price", "sum")])
    for o in orders:
        date = o.created_at.astimezone(orders_service.TIMEZONE).replace(tzinfo=None) if o.created_at.tzinfo else o.created_at
        status = status_text(o.status, lang).split(" —")[0].split(" ", 1)[-1]  # без эмодзи — удобно фильтровать
        u = o.user
        ws.append([
            o.id, date, status, person_name(u), u.phone,
            localized_name(u.region, lang) if u.region else None, o.dealer.name,
            localized_name(o.dealer.region, lang), sum(i.quantity for i in o.items), float(o.total_amount),
        ])
        for i in o.items:
            items_ws.append([o.id, date, status, person_name(u), u.phone, o.dealer.name, localized_name(i, lang),
                             i.part_number, i.model_name, i.node_name, i.quantity, float(i.price), float(i.total)])
    for sheet, date_col, money_cols in ((ws, 2, (10,)), (items_ws, 2, (12, 13))):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in sheet.iter_rows(min_row=2):
            row[date_col - 1].number_format = "DD.MM.YYYY HH:MM"
            for c in money_cols:
                row[c - 1].number_format = "#,##0"
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 18
        sheet.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_filename(period: str) -> str:
    from datetime import datetime
    return f"jac_orders_{period}_{datetime.now(orders_service.TIMEZONE):%Y-%m-%d}.xlsx"


__all__ = ["stats_text", "orders_excel", "export_filename", "fmt_date"]
